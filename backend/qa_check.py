"""Comprehensive end-to-end QA against a RUNNING server.

Usage: python qa_check.py [base_url]   (default http://127.0.0.1:8123)
Exercises every user-facing flow over HTTP and prints a PASS/FAIL report.
"""

import io
import sys
import time
import zipfile

import httpx
import openpyxl

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8123"
results: list[tuple[str, bool, str]] = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(f"  {'PASS' if cond else 'FAIL'}  {name}" + (f" — {detail}" if detail else ""))


def wait_job(client, token, job_id, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        r = client.get(f"/api/jobs/{job_id}", headers={"X-Session-Token": token})
        rec = r.json()
        if rec["status"] in ("completed", "failed"):
            return rec
        time.sleep(0.1)
    return {"status": "timeout"}


def xlsx_bytes(sheets):
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xls_bytes(sheets):
    import xlwt
    wb = xlwt.Workbook()
    for name, rows in sheets.items():
        ws = wb.add_sheet(name)
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                ws.write(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    c = httpx.Client(base_url=BASE, timeout=60)

    print("\n== Health & session ==")
    check("health endpoint", c.get("/api/health").status_code == 200)
    check("unauthenticated request rejected", c.get("/api/project").status_code == 401)
    token = c.post("/api/sessions").json()["token"]
    check("session token issued", bool(token))
    H = {"X-Session-Token": token}

    print("\n== CSV upload (dirty data) ==")
    dirty = (
        " name ,Age,Dept,Salary,Joined,Notes\n"
        " Alice ,30,Eng,100,2024-01-01,ok\n"
        "bob,-5,eng,100,01/02/2024,ok\n"
        "Alice,30,Eng,100,2024-01-01,ok\n"
        "N/A,40,Sales,200,2024-01-04,ok\n"
        "CAROL,45,sales,200,2024-01-05,ok\n"
        "dave,N/A,Eng,100,bad,ok\n"
        "eve,50,,300,2024-01-07,ok\n"
        "Alice,30,Eng,100,2024-01-01,ok\n"
        "=formula,25,Eng,99,2024-01-08,ok\n"
    )
    r = c.post("/api/upload", files={"file": ("messy.csv", dirty.encode(), "text/csv")}, headers=H)
    check("upload accepted (202)", r.status_code == 202, str(r.status_code))
    job = wait_job(c, token, r.json()["job_id"])
    check("ingest job completed", job["status"] == "completed", job.get("error", ""))
    project = c.get("/api/project", headers=H).json()
    iid = project["items"][0]["item_id"]
    check("project has 1 csv item", len(project["items"]) == 1)

    r = c.get(f"/api/items/{iid}/preview", headers=H)
    check("preview returns rows+dtypes", r.status_code == 200 and len(r.json()["rows"]) > 0)

    print("\n== Cleaning flow ==")
    # missingness
    r = c.get(f"/api/items/{iid}/assess/missingness", headers=H)
    check("missingness assess", r.status_code == 200)
    missing_cols = [x["column"] for x in r.json()["columns"] if x["missing_count"] > 0]
    check("missingness found markers", set(missing_cols) >= {"Dept", "Age", " name "}, str(missing_cols))

    r = c.post(f"/api/items/{iid}/ops", headers=H, json={
        "op_type": "treat_missing", "stage": "missingness",
        "params": {"column": "Dept", "method": "constant", "value": "Unknown"}})
    check("fill missing Dept", r.status_code == 200)

    # rename
    r = c.post(f"/api/items/{iid}/ops", headers=H, json={
        "op_type": "rename_columns", "stage": "column_names",
        "params": {"mapping": {" name ": "name", "Age": "age", "Dept": "dept",
                               "Salary": "salary", "Joined": "joined", "Notes": "notes"}}})
    check("rename columns", r.status_code == 200)

    # types
    r = c.get(f"/api/items/{iid}/assess/types", headers=H)
    check("type assess", r.status_code == 200)
    r = c.post(f"/api/items/{iid}/ops", headers=H, json={
        "op_type": "convert_type", "stage": "types",
        "params": {"column": "age", "target": "integer"}})
    check("convert age->int", r.status_code == 200)

    # normalize
    r = c.post(f"/api/items/{iid}/ops", headers=H, json={
        "op_type": "normalize", "stage": "normalization",
        "params": {"columns": ["name"], "operations": [{"op": "trim"}, {"op": "lower"}]}})
    check("normalize name", r.status_code == 200)

    # invalid values
    r = c.post(f"/api/items/{iid}/ops", headers=H, json={
        "op_type": "mask_values", "stage": "invalid_values",
        "params": {"column": "age", "rule": {"type": "range", "min": 0, "max": 120}}})
    check("mask invalid age", r.status_code == 200)
    prev = c.get(f"/api/items/{iid}/preview?n=50", headers=H).json()
    ages = [row["age"] for row in prev["rows"]]
    check("-5 masked to null", ages[1] is None)

    # duplicates
    r = c.post(f"/api/items/{iid}/duplicates/estimate", headers=H,
               json={"columns": ["name", "salary"], "blocking_columns": [], "threshold": 100})
    check("dup estimate", r.status_code == 200)
    r = c.post(f"/api/items/{iid}/duplicates/analyze", headers=H,
               json={"columns": ["name", "salary"], "blocking_columns": [], "threshold": 100})
    check("dup analyze job", r.status_code == 202)
    job = wait_job(c, token, r.json()["job_id"])
    check("dup analyze completed", job["status"] == "completed", job.get("error", ""))
    dres = c.get(f"/api/items/{iid}/duplicates/result", headers=H).json()
    check("dup groups found", dres["summary"]["duplicate_groups"] >= 1, str(dres["summary"]))
    gid = dres["groups"][0]["members"]
    r = c.post(f"/api/items/{iid}/duplicates/resolve", headers=H,
               json={"decisions": [{"row_ids": gid, "action": "keep_first"}]})
    check("dup resolve keep_first", r.status_code == 200, str(r.status_code))
    prev = c.get(f"/api/items/{iid}/preview?n=50", headers=H).json()
    check("dup rows removed", prev["row_count"] == 7, str(prev["row_count"]))

    # outliers
    r = c.get(f"/api/items/{iid}/outliers/screen/salary?method=manual&min=0&max=250", headers=H)
    check("outlier screen", r.status_code == 200)
    flagged = [f["row_id"] for f in r.json()["flagged"]]
    check("outlier flagged", len(flagged) >= 1, str(flagged))
    r = c.post(f"/api/items/{iid}/outliers/treat", headers=H,
               json={"column": "salary", "row_ids": flagged, "action": "remove"})
    check("treat gated pre-review (409)", r.status_code == 409)
    r = c.post(f"/api/items/{iid}/outliers/bivariate", headers=H,
               json={"column": "salary", "context": "dept"})
    check("bivariate review", r.status_code == 200, str(r.status_code))
    r = c.post(f"/api/items/{iid}/outliers/treat", headers=H,
               json={"column": "salary", "row_ids": flagged, "action": "keep"})
    check("keep after review", r.status_code == 200)

    print("\n== Undo/redo ==")
    r = c.get(f"/api/items/{iid}/history", headers=H).json()
    check("history lists ops", r["pointer"] > 0)
    r = c.post(f"/api/items/{iid}/undo", headers=H)
    check("undo", r.status_code == 200)
    r = c.post(f"/api/items/{iid}/redo", headers=H)
    check("redo", r.status_code == 200)

    print("\n== Validation & export ==")
    r = c.post(f"/api/items/{iid}/validate", headers=H)
    job = wait_job(c, token, r.json()["job_id"])
    check("validation job", job["status"] == "completed", job.get("error", ""))
    v = job["result"]
    check("validation has before/after", "row_count" in v and "cells_changed" in v)

    w = c.get("/api/export/warnings", headers=H).json()
    print(f"    warnings: {w['warnings']}")
    r = c.post("/api/export", headers=H,
               json={"formats": ["csv", "xlsx", "report", "json", "script", "flagged", "zip"],
                     "acknowledge_warnings": True})
    job = wait_job(c, token, r.json()["job_id"], timeout=60)
    check("export job completed", job["status"] == "completed", job.get("error", ""))
    files = {f["name"] for f in job["result"]["files"]}
    expected = {"messy_cleaned.csv", "messy_cleaned.xlsx", "messy_report.html",
                "messy_transformations.json", "messy_clean.py", "messy_cleaned.zip"}
    check("all export artifacts produced", expected <= files, str(files - expected))

    for name in expected - {"messy_cleaned.zip"}:
        r = c.get(f"/api/exports/{name}", headers=H)
        check(f"download {name}", r.status_code == 200 and len(r.content) > 0)
    body = c.get("/api/exports/messy_cleaned.csv", headers=H).text
    check("formula injection escaped", "'=formula" in body)
    zf = zipfile.ZipFile(io.BytesIO(c.get("/api/exports/messy_cleaned.zip", headers=H).content))
    check("zip contains artifacts", len(zf.namelist()) >= 5, str(zf.namelist()))

    print("\n== Multi-sheet XLSX ==")
    t2 = c.post("/api/sessions").json()["token"]
    data = xlsx_bytes({"Main": [["a", "b"], [1, 2], [3, 4]], "Aux": [["x"], [9]]})
    r = c.post("/api/upload", files={"file": ("book.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
               headers={"X-Session-Token": t2})
    job = wait_job(c, t2, r.json()["job_id"])
    check("xlsx upload", job["status"] == "completed", job.get("error", ""))
    proj = c.get("/api/project", headers={"X-Session-Token": t2}).json()
    check("xlsx has 2 items", len(proj["items"]) == 2)

    print("\n== Legacy .xls ==")
    data = xls_bytes({"S1": [["a"], [1], [2]]})
    r = c.post("/api/upload", files={"file": ("old.xls", data, "application/vnd.ms-excel")},
               headers={"X-Session-Token": t2})
    job = wait_job(c, t2, r.json()["job_id"])
    check("xls upload parses", job["status"] == "completed", job.get("error", ""))
    proj = c.get("/api/project", headers={"X-Session-Token": t2}).json()
    check("xls project loaded", proj.get("kind") == "xls" and len(proj["items"]) == 1)

    print("\n== Rejections & security ==")
    r = c.post("/api/upload", files={"file": ("evil.exe", b"MZ\x90\x00", "application/octet-stream")},
               headers={"X-Session-Token": t2})
    check("bad extension rejected", r.status_code == 422)
    r = c.post("/api/upload", files={"file": ("fake.csv", b"not really", "text/csv")},
               headers={"X-Session-Token": "bad-token-123"})
    check("forged token rejected", r.status_code == 401)
    r = c.get("/api/exports/messy_cleaned.csv", headers={"X-Session-Token": t2})
    check("cross-session export denied", r.status_code == 404)
    r = c.get("/api/exports/../meta.json", headers={"X-Session-Token": t2})
    check("path traversal denied", r.status_code in (404, 422))

    print("\n== Corrupted/garbage file ==")
    r = c.post("/api/upload", files={"file": ("bad.xlsx", b"\x89PNGgarbage", "application/vnd.ms-excel")},
               headers={"X-Session-Token": t2})
    job = wait_job(c, t2, r.json()["job_id"])
    check("corrupted xlsx fails cleanly", job["status"] == "failed")

    print("\n== Frontend serving ==")
    r = c.get("/")
    check("SPA index served", r.status_code == 200 and "text/html" in r.headers.get("content-type", ""))

    failed = [x for x in results if not x[1]]
    print(f"\n{'='*50}\n{len(results) - len(failed)}/{len(results)} checks passed"
          + (f" — {len(failed)} FAILED" if failed else ""))
    for name, _, detail in failed:
        print(f"  FAILED: {name} {detail}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
